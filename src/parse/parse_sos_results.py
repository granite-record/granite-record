#!/usr/bin/env python3
# GRANITE_VERSION: 2026-10-17.2
"""
The Secretary of State's count of one general election's State Representative
and State Senate races, read into collected/election_<year>.json.

    python3 src/parse/parse_sos_results.py --year 2024          # collected/election_2024.json
    python3 src/parse/parse_sos_results.py --year 2024 --check  # compare again; write nothing

WHAT IT READS. The spreadsheets the Secretary of State publishes after a
general election, saved by a person in a browser, byte for byte and under the
names the Secretary of State gave them, in records/sources/ (the person's
decision of 16 October 2026: no fetcher until the 2026 page has been seen):

    <year>-ge-house-<county>[_N].xls or .xlsx      one per county, ten in all
    <year>-ge-state-senate-district-1-24[_N].xls   every Senate district

and site/districts.json, the district table every page draws from, which the
files are held against. It needs openpyxl for .xlsx and xlrd for .xls, and
nothing else outside the standard library. Nothing here asks the network.

WHAT IT WRITES. collected/election_<year>.json, in the shape
src/parse/build_members_elect.py reads (its docstring has the shape): every
race the district table names, with its seats, whether it is floterial, its
towns as the district table spells them, every candidate with a party and a
count of votes, who was elected, and the scattered write-ins. "_source" names
each file read with its sha256 and the day it was saved. A person runs this
and commits the file, so that the diff is the review. With --check it writes
nothing, and says where what it reads now differs from what is committed.

THE LAYOUT, AS THE 2024 FILES HAVE IT, read on 17 October 2026 before a line
of this was written:

  A county's House file is one sheet. Three title rows: the election, "State
  Representative(s) - <COUNTY> County", and the election's day. Then each
  district: a header row "District No. 5 (4)" -- the number, the seats, and
  " F" or "FL" for a floterial -- with a column per candidate, "Wendy Chase,
  d", then Undervotes, Overvotes and Write-Ins; a row per town or ward; and a
  Totals row where there is more than one. A district with more candidates
  than fit across the sheet goes on below, under a header row whose first
  cell is empty, with its own rows per town. A recount shows two ways: a
  "RECOUNT" column beside each candidate's count, which holds the recount's
  figure (Hillsborough 9, 17, 18 and 24; Grafton 7; Sullivan 5), or a second
  whole district, "District No. 4 (3) RECOUNT FIGURES" (Strafford 4).

  The Senate file is a sheet for one to three districts: "State Senate
  District 10", then a header row whose first cell is the election's day and
  whose others name the candidates as the House files do, a row per town or
  ward, and Totals.

  A candidate written in and named is "WRITE-IN   Kathy DesRoches"
  (Merrimack 14), with no party. A party is its letters; a cross-nominated
  candidate's are joined with "&" ("Janet Larsen, i&r"), and are written here
  joined with "/" as build_members_elect reads them.

WINNERS ARE THE TOP OF THE COUNT. The files mark no winner, so the candidates
with the most votes, as many as the race has seats, are elected; where the
last seat's count is shared, those sharing it are neither ("elected": null),
which build_members_elect draws as not final. A race whose count is a
recount's carries "recounted": true.

IT REFUSES RATHER THAN GUESSES, saying which file, sheet and row, and writes
nothing: a title, header or cell it does not recognise; a district or a place
the district table does not name; seats or a district's towns that disagree
with the district table; a race the table names and no file holds; a town's
count missing for a candidate; a column that does not add up to the Totals
the file prints beside it, which are the Secretary of State's own sums; a
write-in the file does not name who may have won -- fewer named candidates
than seats, or scattered write-ins as many as the last seat's votes; and a
file of another election, or a county's file twice under one number.

THE PERSON'S DECISIONS OF 17 OCTOBER 2026 about the reader, each as
recommended after the 2024 files were read: a printed floterial mark the
towns contradict is flagged in "_notes" and decides nothing (a typo in the
official record flags rather than fails); a write-in winner the file does not
name is refused, never guessed; and a county's file saved again after a
recount, under the Secretary of State's next number ("_8" after "_7"),
replaces the earlier, which is said in "_notes" and not read.
"""

# The bootstrap: _paths.py, found above this file, puts every code folder on the import path.
import sys
from pathlib import Path
sys.path += [str(p) for p in Path(__file__).resolve().parents if (p / "_paths.py").is_file()][:1]
import _paths  # noqa: E402,F401

import argparse
import datetime
import hashlib
import json
import re
import unicodedata
from collections import Counter

import build_members_elect as BME

SOURCES = Path("records") / "sources"
COUNTIES = ("belknap", "carroll", "cheshire", "coos", "grafton", "hillsborough", "merrimack",
            "rockingham", "strafford", "sullivan")

HOUSE_FILE = re.compile(r"^(\d{4})-ge-house-([a-z]+)(?:_\d+)?\.xlsx?$")
SENATE_FILE = re.compile(r"^(\d{4})-ge-state-senate-district-1-24(?:_\d+)?\.xlsx?$")

# EVERY SPELLING OF A DISTRICT'S HEADER THE 2024 FILES USE (17 October
# 2026), and no other: "District No. 5 (4)", "District No 17 (2)", "District
# No.7 (1)", "District 16 (1)", "District No. 31 (2)FL", "District No. 8 (2)
# F", and the recount's "District No. 4 (3) RECOUNT FIGURES". A header that
# moves from these is refused, not guessed at.
DISTRICT = re.compile(r"^district(?:\s+no\b\.?)?\s*(\d+)\s*\((\d+)\)\s*(fl|f)?\s*"
                      r"(recount figures)?$", re.I)
SENATE = re.compile(r"^state senate district\s+(\d+)$", re.I)
REP_TITLE = re.compile(r"^state representatives?\s*-\s*([a-z]+)\s+county$", re.I)
# A CANDIDATE IS A NAME, A COMMA AND THE PARTY'S LETTERS (17 October 2026):
# "Mark Proulx,r" has no space, and "Raymond C. Peeples Jr., r" has a comma
# of its own before it, so the last comma is the one that counts.
CANDIDATE = re.compile(r"^(?P<name>\S.*?)\s*,\s*(?P<party>[a-z]{1,3}(?:\s*&\s*[a-z]{1,3})*)$", re.I)
WRITE_IN = re.compile(r"^write[- ]?in\s+(?P<name>\S.*)$", re.I)
META = {"undervotes": "under", "overvotes": "over", "write-ins": "scatter"}
TOTALS = "totals"

# THE UNINCORPORATED PLACES, WHICH EACH FILE ABBREVIATES ITS OWN WAY (17
# October 2026): "At. & Gilm. Ac. Grant", "Atkinson and Gilmanton Ac. Gt." and
# the district table's "Atk. & Gil. Academy Grant" are one place. A word is
# expanded only whole, and the result must then be a place the table names
# exactly, or the file is refused.
ABBREVIATIONS = {"pur": "purchase", "gt": "grant", "loc": "location", "ac": "academy",
                 "at": "atkinson", "atk": "atkinson", "gil": "gilmanton", "gilm": "gilmanton",
                 "thomp": "thompson"}


class Refused(Exception):
    """A file this does not recognise: nothing is written."""


class Unreadable(Refused):
    """A file this machine has no library to open: not a fault in the file."""


# ---- cells -------------------------------------------------------------------

def cell(v):
    """A cell as this reads it: None for an empty one or one of spaces, a
    whole number for a count (xlrd hands every number back as a float), a
    date for a day, and a string with its ends trimmed."""
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, str):
        return " ".join(v.split()) or None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def sheets(path):
    """[(sheet name, [[cell, ...], ...]), ...] of a saved file, by its kind."""
    path = Path(path)
    try:
        if path.suffix.lower() == ".xlsx":
            import openpyxl
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                return [(ws.title, [[cell(v) for v in row]
                                    for row in ws.iter_rows(values_only=True)])
                        for ws in book.worksheets]
            finally:
                book.close()
        import xlrd
    except ImportError as e:
        raise Unreadable(f"{path.name} needs {e.name} to be read: pip install openpyxl==3.1.5 "
                         "xlrd==2.0.2")
    book = xlrd.open_workbook(str(path))
    out = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                x = sh.cell(r, c)
                row.append(xlrd.xldate_as_datetime(x.value, book.datemode).date()
                           if x.ctype == xlrd.XL_CELL_DATE else cell(x.value))
            rows.append(row)
        out.append((sh.name, rows))
    return out


def as_day(v):
    """A day from a date cell, or from the serial number an .xls cell keeps
    a date as when it is not formatted as one (45601 is 5 November 2024)."""
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, int) and 30000 < v < 80000:
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=v)
    return None


def fold(s):
    return unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode("ascii")


def edits(a, b):
    """The fewest letters added, taken away or changed to make `a` into `b`."""
    row = list(range(len(b) + 1))
    for i, x in enumerate(a, 1):
        prev, row[0] = row[0], i
        for j, y in enumerate(b, 1):
            prev, row[j] = row[j], min(row[j] + 1, row[j - 1] + 1, prev + (x != y))
    return row[-1]


def place_key(s):
    """A place's name with its abbreviations spelled out, for matching only.
    fold() has already dropped a curly apostrophe; a straight one and a
    point go here, so that "Mes's" and "Ac." are read as one word."""
    s = re.sub(r"['.]", "", fold(s).lower().replace("&", " and "))
    return " ".join(str(int(w)) if w.isdigit() else ABBREVIATIONS.get(w, w)
                    for w in re.findall(r"[a-z0-9]+", s))


# ---- one block: a header row and the rows under it ------------------------

def entities(row, start, where):
    """The columns a header row names: candidates and the three counts that
    are not a candidate's, each with the column of its recount where a
    "RECOUNT" column follows it."""
    out = []
    for c in range(start, len(row)):
        v = row[c]
        if v is None:
            continue
        if not isinstance(v, str):
            raise Refused(f"{where}: {v!r} in column {c + 1}, where a candidate's name belongs")
        low = v.lower()
        if low == "recount":
            # A RECOUNT COLUMN BELONGS TO THE ONE BESIDE IT (17 October 2026),
            # as every one in the 2024 files does, and only to that one:
            # anywhere else it would be a guess which count it revises.
            if not out or out[-1]["recount"] is not None or out[-1]["col"] != c - 1:
                raise Refused(f"{where}: a RECOUNT column in column {c + 1} follows no count")
            out[-1]["recount"] = c
            continue
        if low in META:
            out.append({"meta": META[low], "col": c, "recount": None})
            continue
        m = WRITE_IN.match(v)
        if m:
            out.append({"name": m.group("name").strip(), "party": "", "write_in": True,
                        "col": c, "recount": None})
            continue
        m = CANDIDATE.match(v)
        if not m:
            raise Refused(f"{where}: {v!r} in column {c + 1} is neither a candidate (\"Name, d\") "
                          "nor Undervotes, Overvotes or Write-Ins")
        party = "/".join(p.strip().lower() for p in m.group("party").split("&"))
        out.append({"name": m.group("name").strip(), "party": party, "write_in": False,
                    "col": c, "recount": None})
    if not any("name" in e for e in out):
        raise Refused(f"{where}: a header row that names no candidate")
    return out


class Block:
    """A header row and the rows under it, to the next header."""

    def __init__(self, ents, where, recount_figures=False):
        self.ents, self.where, self.rows, self.totals = ents, where, [], None
        self.recount_figures = recount_figures
        # (entity, 0 for the count or 1 for its recount) that a Totals row held.
        self.held = set()

    def count(self, row, col, meta, where):
        v = row[col] if col < len(row) else None
        if v is None and meta:
            return None
        if not isinstance(v, int) or isinstance(v, bool) or v < 0:
            raise Refused(f"{where}: {v!r} in column {col + 1}, where a count belongs")
        return v

    def read(self, single_town):
        """{town: {entity index: (count, recount)}}, each column held to the
        Totals row the file prints, where it prints one."""
        out = {}
        for town, row, where in self.rows:
            if town is None:
                # A ROW WITH NO TOWN IS THE DISTRICT'S ONE TOWN (17 October
                # 2026), as under Hillsborough 9's second header in 2024;
                # where the district has more than one, which it is would be
                # a guess.
                if single_town is None:
                    raise Refused(f"{where}: a row of counts that names no town, in a district "
                                  "of more than one")
                town = single_town
            if town in out:
                raise Refused(f"{where}: {town} twice in one block")
            out[town] = {i: (self.count(row, e["col"], "meta" in e, where),
                             None if e["recount"] is None
                             else self.count(row, e["recount"], "meta" in e, where))
                         for i, e in enumerate(self.ents)}
        if not out:
            raise Refused(f"{self.where}: a header with no rows of counts under it")
        if self.totals is None:
            if len(out) > 1:
                raise Refused(f"{self.where}: {len(out)} towns and no Totals row to hold them to")
            return out
        row, where = self.totals
        for i, e in enumerate(self.ents):
            for k, col in ((0, e["col"]), (1, e["recount"])):
                if col is None:
                    continue
                said = self.count(row, col, "meta" in e, where)
                if said is None:
                    continue
                # AN EMPTY CELL THE TOTALS ACCOUNT FOR IS NOUGHT (17 October
                # 2026). Senate 3's Madison prints no write-ins in 2024 and
                # the district's Totals say 19, which the other towns make
                # without it.
                self.held.add((i, k))
                got = sum(v[i][k] or 0 for v in out.values())
                if got != said:
                    raise Refused(f"{where}: {e.get('name') or e['meta']}'s "
                                  f"{'recount ' if k else ''}column adds up to {got:,} and the "
                                  f"file's Totals say {said:,}")
        return out


# ---- a district, from its blocks -------------------------------------------

def district(blocks, seats, where):
    """{"towns", "candidates", "scatter", "recounted"} from a district's
    blocks: the first and any under it, and a whole recount's block, whose
    figures replace the first count's."""
    first = [b for b in blocks if not b.recount_figures]
    again = [b for b in blocks if b.recount_figures]
    if len(again) > 1:
        raise Refused(f"{where}: more than one block of recount figures")

    def counts(group):
        # `nought`: the towns whose empty cell of a count a Totals row
        # accounted for, which is then nought where no block prints it.
        cands, meta, towns, nought = {}, {}, None, {}
        recounted = False
        single = None
        for b in group:
            got = b.read(single)
            if towns is None:
                towns = list(got)
                single = towns[0] if len(towns) == 1 else None
            elif set(got) != set(towns):
                raise Refused(f"{b.where}: the towns {sorted(got)} under this header are not the "
                              f"district's {sorted(towns)}")
            for i, e in enumerate(b.ents):
                if e["recount"] is not None:
                    recounted = True
                if "meta" in e:
                    k = 1 if e["recount"] is not None else 0
                    if (i, k) in b.held:
                        nought.setdefault(e["meta"], set()).update(got)
                    for t, v in got.items():
                        n = v[i][k]
                        if n is None:
                            continue
                        old = meta.setdefault(e["meta"], {}).setdefault(t, n)
                        if old != n:
                            raise Refused(f"{b.where}: {t}'s {e['meta']} is {old} in one block "
                                          f"and {n} in another")
                    continue
                if e["name"] in cands:
                    raise Refused(f"{b.where}: {e['name']} twice in one district")
                votes = sum(v[i][1] if e["recount"] is not None else v[i][0]
                            for v in got.values())
                cands[e["name"]] = {"name": e["name"], "party": e["party"], "votes": votes,
                                    "write_in": e["write_in"]}
        scatter = meta.get("scatter") or {}
        for t in nought.get("scatter") or ():
            scatter.setdefault(t, 0)
        return (towns, list(cands.values()),
                sum(scatter.values()) if set(scatter) == set(towns) else None, recounted)

    towns, cands, scatter, recounted = counts(first)
    if again:
        t2, c2, s2, _ = counts(again)
        if set(t2) != set(towns) or {c["name"] for c in c2} != {c["name"] for c in cands}:
            raise Refused(f"{again[0].where}: the recount's towns or candidates are not the "
                          "first count's")
        cands, scatter, recounted = c2, s2, True
    return {"towns": towns, "candidates": elect(cands, seats, where, scatter),
            "scatter": scatter, "recounted": recounted}


def elect(cands, seats, where, scatter=None):
    """The candidates, each marked elected, not, or tied for the last seat
    (null): the files mark no winner, so the count decides. `scatter` is
    the write-ins the file counts and does not name, or None where it does
    not count them for every place."""
    if len(cands) < seats:
        raise Refused(f"{where}: {len(cands)} named candidate(s) for {seats} seat(s) -- a write-in "
                      "the file does not name may have won, and this does not guess who")
    by = sorted((c["votes"] for c in cands), reverse=True)
    cut = by[seats - 1]
    if cut == 0:
        raise Refused(f"{where}: a seat whose winner has no votes")
    # A WRITE-IN WINNER THE FILE DOES NOT NAME IS REFUSED, NOT GUESSED (the
    # person's decision of 17 October 2026). The scattered write-ins are one
    # count for many names, so where it reaches the last seat's votes one of
    # those names may hold the seat, and which is not in the file. In 2024
    # the most it came to was 128 against 3,389 (Rockingham 4).
    if scatter is not None and scatter >= cut:
        raise Refused(f"{where}: the write-ins the file does not name come to {scatter:,}, and the "
                      f"last seat was won with {cut:,} -- one of them may have won, and this does "
                      "not guess who")
    above = sum(1 for v in by if v > cut)
    at = sum(1 for v in by if v == cut)
    tie = above + at > seats
    return [dict(c, elected=(True if c["votes"] > cut or (c["votes"] == cut and not tie)
                             else None if c["votes"] == cut else False)) for c in cands]


# ---- the files ---------------------------------------------------------------

def blank(row):
    return all(v is None for v in row)


def read_house(path, county, election):
    """{district number: (seats, floterial, district record)} from one
    county's House file."""
    got = sheets(path)
    if len(got) != 1:
        raise Refused(f"{path.name}: {len(got)} sheets where a county's House file has one")
    name, rows = got[0]

    def at(i):
        return f"{path.name}, sheet {name!r}, row {i + 1}"
    title = " ".join(str(c) for c in (rows[0] if rows else []) if c)
    if str(election.year) not in title or "general election" not in title.lower():
        raise Refused(f"{at(0)}: {title!r} is not the title of the {election.year} general "
                      "election")
    m = REP_TITLE.match(str((rows[1] if len(rows) > 1 else [None])[0] or ""))
    if not m or m.group(1).lower() != county:
        raise Refused(f"{at(1)}: {rows[1][:1] if len(rows) > 1 else []!r} is not \"State "
                      f"Representatives - {county.upper()} County\"")
    day = as_day((rows[2] if len(rows) > 2 else [None])[0])
    if day != election:
        raise Refused(f"{at(2)}: the file's day is {day}, not the election of {election}")
    found, cur = {}, None
    for i in range(3, len(rows)):
        row = rows[i]
        while row and row[-1] is None:
            row = row[:-1]
        if blank(row):
            continue
        head = row[0]
        m = DISTRICT.match(head) if isinstance(head, str) else None
        if m:
            n, seats = int(m.group(1)), int(m.group(2))
            flo, again = bool(m.group(3)), bool(m.group(4))
            if again != (n in found):
                raise Refused(f"{at(i)}: district {n} " + ("recounted before it is counted"
                                                           if again else "twice"))
            if not again:
                found[n] = (seats, flo, [], at(i))
            elif (seats, flo) != found[n][:2]:
                raise Refused(f"{at(i)}: the recount's seats or floterial mark are not the "
                              "first count's")
            cur = found[n]
            cur[2].append(Block(entities(row, 1, at(i)), at(i), recount_figures=again))
            continue
        if cur is None:
            raise Refused(f"{at(i)}: {row!r} before any district's header")
        if head is None and any(isinstance(v, str) for v in row[1:]):
            # THE COUNTS THAT ARE NOT A CANDIDATE'S KEEP THEIR COLUMNS (17
            # October 2026) under a second header that does not name them
            # again: in 2024 Rockingham 17's Windham prints its undervotes,
            # overvotes and write-ins on the row under the second header, in
            # the columns the first header named.
            ents = entities(row, 1, at(i))
            taken = {c for e in ents for c in (e["col"], e["recount"]) if c is not None}
            named = {e["meta"] for e in ents if "meta" in e}
            ents += [dict(e) for e in cur[2][0].ents if "meta" in e and e["meta"] not in named
                     and not {e["col"], e["recount"]} & taken]
            cur[2].append(Block(ents, at(i), recount_figures=cur[2][-1].recount_figures))
            continue
        if isinstance(head, str) and head.lower() == TOTALS:
            if cur[2][-1].totals is not None:
                raise Refused(f"{at(i)}: a second Totals row under one header")
            cur[2][-1].totals = (row, at(i))
            continue
        if head is not None and not isinstance(head, str):
            raise Refused(f"{at(i)}: {head!r} where a town's name belongs")
        if cur[2][-1].totals is not None:
            raise Refused(f"{at(i)}: a row of counts after the Totals")
        cur[2][-1].rows.append((head, row, at(i)))
    if not found:
        raise Refused(f"{path.name}: no district in it")
    return {n: (seats, flo, district(blocks, seats, where))
            for n, (seats, flo, blocks, where) in found.items()}


def read_senate(path, election):
    """{district number: district record} from the Senate file."""
    found, cur = {}, None
    for name, rows in sheets(path):
        cur = None
        for i, row in enumerate(rows):
            where = f"{path.name}, sheet {name!r}, row {i + 1}"
            while row and row[-1] is None:
                row = row[:-1]
            if blank(row):
                continue
            head = row[0]
            m = SENATE.match(head) if isinstance(head, str) else None
            if m:
                n = int(m.group(1))
                if n in found:
                    raise Refused(f"{where}: Senate district {n} twice")
                cur = found[n] = {"n": n, "where": where, "blocks": []}
                continue
            if isinstance(head, str) and "general election" in head.lower() and len(row) == 1:
                continue
            if cur is None:
                raise Refused(f"{where}: {row!r} before any district's title")
            if not cur["blocks"]:
                day = as_day(head)
                if day != election:
                    raise Refused(f"{where}: the header's day is {head!r}, not the election of "
                                  f"{election}")
                cur["blocks"].append(Block(entities(row, 1, where), where))
                continue
            b = cur["blocks"][-1]
            if isinstance(head, str) and head.lower() == TOTALS:
                if b.totals is not None:
                    raise Refused(f"{where}: a second Totals row")
                b.totals = (row, where)
                continue
            if not isinstance(head, str):
                raise Refused(f"{where}: {head!r} where a town's name belongs")
            if b.totals is not None:
                raise Refused(f"{where}: a row of counts after the Totals")
            b.rows.append((head, row, where))
    if not found:
        raise Refused(f"{path.name}: no Senate district in it")
    for n, d in found.items():
        if not d["blocks"]:
            raise Refused(f"{d['where']}: Senate district {n} has no header row")
    return {n: district(d["blocks"], 1, d["where"]) for n, d in found.items()}


def upload_number(p):
    """Where a saved file stands among the Secretary of State's uploads of
    one name: none (-1) before "_0", "_0" before "_1", and so on."""
    m = re.search(r"_(\d+)\.xlsx?$", p.name)
    return int(m.group(1)) if m else -1


def latest(files, what, replaced):
    """The one of `files` the Secretary of State uploaded last, by the
    number on its name; each earlier one is added to `replaced` as (the
    earlier, the later).

    A RECOUNT'S RE-DOWNLOAD REPLACES THE COUNTY'S FILE (the person's
    decision of 17 October 2026). The number on a saved file's name rises
    with each upload of that name -- the 2024 Hillsborough file, which
    carries four recounts, is "_7" -- so a recount's figures saved beside
    the first count are, by that sign and not by anything the Secretary of
    State has said, the higher number; and the later is the results. Until
    then two files of one county stopped the reader. Two under the same
    number are still refused: which is the later is then not in their
    names."""
    files = sorted(files, key=lambda p: (upload_number(p), p.name))
    if len(files) > 1 and upload_number(files[-1]) == upload_number(files[-2]):
        raise Refused(f"{files[-2].name} and {files[-1].name} are both {what}, under one number: "
                      "which the Secretary of State published later is not in their names; keep "
                      "the one saved last")
    replaced.extend((old, files[-1]) for old in files[:-1])
    return files[-1]


def source_files(root, year):
    """{"house": {county: path}, "senate": path, "replaced": [(earlier,
    later), ...]}: the year's saved files, each county's latest once."""
    d = Path(root) / SOURCES
    house, senate, replaced = {}, [], []
    for p in sorted(d.glob(f"{year}-ge-*")) if d.is_dir() else []:
        m = HOUSE_FILE.match(p.name)
        if m and m.group(1) == str(year):
            if m.group(2) not in COUNTIES:
                raise Refused(f"{p.name} names no county of New Hampshire")
            house.setdefault(m.group(2), []).append(p)
        elif SENATE_FILE.match(p.name):
            senate.append(p)
    missing = [c for c in COUNTIES if c not in house]
    if missing:
        raise Refused(f"no {year} House results saved in {d.as_posix()}/ for "
                      + ", ".join(missing))
    if not senate:
        raise Refused(f"no {year} State Senate file in {d.as_posix()}/")
    return {"house": {c: latest(ps, f"{c}'s House results", replaced) for c, ps in house.items()},
            "senate": latest(senate, "the State Senate results", replaced),
            "replaced": replaced}


def described(p, root, saved=None):
    """A _source entry: the file, as a path from the project's root, its
    sha256, and the day it was saved -- its modification day, unless stated."""
    data = p.read_bytes()
    day = saved or datetime.date.fromtimestamp(p.stat().st_mtime).isoformat()
    return {"file": p.relative_to(root).as_posix(), "url": None, "saved": day,
            "sha256": hashlib.sha256(data).hexdigest()}


# ---- the whole election -------------------------------------------------------

def read_all(root, year, districts, saved=None):
    """The results document for `year`, from the saved files under `root`
    and the district table `districts` (site/districts.json's contents).
    Refused on anything it does not recognise."""
    root = Path(root).resolve()
    election = BME.general_election(year)
    files = source_files(root, year)
    on_file, clash = BME.races_on_file(districts)
    if clash:
        raise Refused("the district table names districts that are no race: "
                      + "; ".join(clash[:4]))
    places = {}
    for r in on_file.values():
        for t in r["towns"]:
            k = place_key(t)
            if places.setdefault(k, t) != t:
                raise Refused(f"the district table names {places[k]!r} and {t!r}, which read as "
                              "one place")
    races = {}
    notes = [f"{new.name} is read, and replaces {old.name}: the Secretary of State's later upload "
             "of the same results, as after a recount" for old, new in files["replaced"]]

    def settle(key, rec, seats, flo, where):
        if key not in on_file:
            raise Refused(f"{where}: {key} is no district in the district table")
        t = on_file[key]
        if seats != t["seats"]:
            raise Refused(f"{where}: {seats} seat(s); the district table says {t['seats']}")
        towns, unknown = [], []
        for town in rec["towns"]:
            p = places.get(place_key(town))
            (towns if p else unknown).append(p or town)
        # ONE MISSPELT PLACE IS READ, AND FLAGGED, WHERE IT CAN BE NOTHING ELSE
        # (17 October 2026). The 2024 Rockingham file names district 10's
        # towns "Newfield" and Newmarket; the table's are Newfields and
        # Newmarket, and the same file spells Newfields right under district
        # 33. A name the table does not hold is taken as the table's only
        # when it is the one place of the district left unmatched and is
        # within two letters of it. Anything more is refused.
        left = sorted(set(t["towns"]) - set(towns))
        if len(unknown) == 1 and len(left) == 1 and \
                edits(place_key(unknown[0]), place_key(left[0])) <= 2:
            notes.append(f"{where}: the file names {unknown[0]!r}, read as {left[0]}, the one "
                         f"place of {key} it names no other way")
            towns.append(left[0])
        elif unknown:
            raise Refused(f"{where}: {', '.join(map(repr, unknown))} "
                          + ("is no place" if len(unknown) == 1 else "are no places")
                          + " the district table names")
        if sorted(towns) != sorted(t["towns"]):
            raise Refused(f"{where}: {key}'s towns are {', '.join(sorted(towns))} here and "
                          f"{', '.join(sorted(t['towns']))} in the district table")
        race = {"seats": seats, "floterial": t["floterial"], "towns": sorted(towns),
                "candidates": rec["candidates"], "scatter": rec["scatter"]}
        if rec["recounted"]:
            race["recounted"] = True
        # WHERE THE FILE COUNTS NO SCATTERED WRITE-INS, THAT IS SAID (17
        # October 2026): elect() could not hold them to the last seat, and
        # silence is not success. In 2024, Hillsborough 12's one town prints
        # none.
        if rec["scatter"] is None:
            notes.append(f"{where}: the file counts no scattered write-ins for every place, so "
                         "whether a write-in it does not name won could not be checked")
        races[key] = race

    for county, path in files["house"].items():
        got = read_house(path, county, election)
        for n, (seats, flo, rec) in got.items():
            settle(f"H|{county.title()}|{n}", rec, seats, flo, f"{path.name}, district {n}")
        # WHICH DISTRICTS FLOAT IS READ FROM THE FILE'S OWN TOWNS, and held to
        # the table (17 October 2026). The printed mark is wrong four times
        # in 2024: Cheshire 14 (Rindge alone) and Grafton 16 (Enfield alone)
        # are marked "FL", and Hillsborough 44 (Goffstown and Weare) and
        # Strafford 18 are not, though each of those two is exactly the towns
        # of districts beneath it. A floterial is a district whose towns are
        # the towns of the other districts that lie inside it, and that is a
        # fact of the file's own rows. Where it and the table disagree the
        # file is refused; where only the printed mark does, the mark is
        # flagged in the file for the review, as a typo in the official record
        # is (the person's rule), and decides nothing.
        towns = {n: set(races[f"H|{county.title()}|{n}"]["towns"]) for n in got}
        for n, (seats, flo, rec) in got.items():
            key, where = f"H|{county.title()}|{n}", f"{path.name}, district {n}"
            inside = [towns[m] for m in towns if m != n and towns[m] <= towns[n]]
            floats = bool(inside) and set().union(*inside) == towns[n]
            if floats != races[key]["floterial"]:
                raise Refused(f"{where}: its towns make it " + ("a floterial" if floats else
                                                                "a base district")
                              + "; the district table says otherwise")
            if flo != floats:
                notes.append(f"{where}: the file " + ("marks it floterial" if flo else
                                                      "does not mark it floterial")
                             + "; by its towns, and in the district table, it is "
                             + ("floterial" if floats else "a base district"))
    for n, rec in read_senate(files["senate"], election).items():
        settle(f"S|{n}", rec, 1, False, f"{files['senate'].name}, district {n}")
    missing = sorted(set(on_file) - set(races), key=BME.race_order)
    if missing:
        raise Refused("the district table names races no file holds: " + ", ".join(missing[:8])
                      + (" ..." if len(missing) > 8 else ""))
    sources = [described(p, root, saved) for p in
               [*(files["house"][c] for c in COUNTIES), files["senate"]]]
    doc = {"_about": [
                "What the Secretary of State's results files say of the State Representative and "
                "State Senate races, read by src/parse/parse_sos_results.py from the files named "
                "in _source and held to the district table. Nothing in it is this project's "
                "decision but the winners, which the files do not mark: the candidates with the "
                "most votes, as many as the race has seats; a tie for the last seat is null.",
                "Write the file again by running the reader, never by hand: --check says where "
                "the saved files and this disagree."],
            "_source": sources,
            "election": election.isoformat(), "term": BME.term_after(year),
            "races": {k: races[k] for k in sorted(races, key=BME.race_order)}}
    if notes:
        doc["_notes"] = notes
    return doc


def dumps(doc):
    """The document as one race to a line or few, so that a recount's diff
    is the race it moved and no more."""
    out = ["{"]
    out.append(' "_about": ' + json.dumps(doc["_about"], ensure_ascii=False, indent=2)
               .replace("\n", "\n ") + ",")
    if doc.get("_notes"):
        out.append(' "_notes": ' + json.dumps(doc["_notes"], ensure_ascii=False, indent=2)
                   .replace("\n", "\n ") + ",")
    out.append(' "_source": [')
    out.append(",\n".join("  " + json.dumps(s, ensure_ascii=False) for s in doc["_source"]))
    out.append(" ],")
    out.append(f' "election": {json.dumps(doc["election"])},')
    out.append(f' "term": {json.dumps(doc["term"])},')
    out.append(' "races": {')
    keys = list(doc["races"])
    for i, k in enumerate(keys):
        r = dict(doc["races"][k])
        cands = r.pop("candidates")
        head = json.dumps(r, ensure_ascii=False)[:-1]
        body = ",\n".join("    " + json.dumps(c, ensure_ascii=False) for c in cands)
        out.append(f'  {json.dumps(k)}: {head}, "candidates": [\n{body}]}}'
                   + ("," if i < len(keys) - 1 else ""))
    out.append(" }")
    out.append("}")
    return "\n".join(out) + "\n"


def differences(now, committed):
    """Where what the files say now and the committed document part, the
    days the files were saved aside: a checkout does not keep a file's
    modification day, so that is the one thing the files cannot say again."""
    def plain(d):
        d = json.loads(json.dumps(d))
        for s in d.get("_source") or []:
            s.pop("saved", None)
        return d
    a, b = plain(now), plain(committed)
    out = []
    for k in ("_notes", "_source", "election", "term"):
        if a.get(k) != b.get(k):
            out.append(f"{k}: the files say {json.dumps(a.get(k))[:200]}, the committed file "
                       f"{json.dumps(b.get(k))[:200]}")
    ra, rb = a.get("races") or {}, b.get("races") or {}
    for k in sorted(set(ra) | set(rb), key=lambda k: BME.race_order(k) if BME.KEY_RE.match(k)
                    else (2, k, 0)):
        if ra.get(k) != rb.get(k):
            out.append(f"{k}: " + ("not in the committed file" if k not in rb else
                                   "not in the files" if k not in ra else
                                   "the files and the committed file differ"))
    return out


def summary(doc):
    races = doc["races"].values()
    cands = [c for r in races for c in r["candidates"]]
    won = Counter(str(c["elected"]) for c in cands)
    recounted = [k for k, r in doc["races"].items() if r.get("recounted")]
    tied = [k for k, r in doc["races"].items() if any(c["elected"] is None
                                                      for c in r["candidates"])]
    house = [k for k in doc["races"] if k.startswith("H|")]
    return (f"{len(doc['_source'])} files; {len(house)} House races and "
            f"{len(doc['races']) - len(house)} Senate, {sum(r['seats'] for r in races)} seats; "
            f"{len(cands)} candidates, {won['True']} elected"
            + (f", {won['None']} tied for a last seat in {', '.join(tied)}" if tied else "")
            + (f"; recounted: {', '.join(recounted)}" if recounted else "")
            + f". Every column agrees with the file's Totals and every race elects its seats.")


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0].strip())
    ap.add_argument("--year", type=int, required=True, help="the general election's year")
    ap.add_argument("--districts", default="site/districts.json")
    ap.add_argument("--out", default=None, help="default collected/election_<year>.json")
    ap.add_argument("--saved", default=None, help="the day the files were saved, YYYY-MM-DD; "
                                                  "default each file's modification day")
    ap.add_argument("--check", action="store_true", help="compare with the committed file; "
                                                         "write nothing")
    a = ap.parse_args()
    out = Path(a.out or f"collected/election_{a.year}.json")
    try:
        if a.year % 2:
            raise Refused(f"{a.year} has no state general election")
        try:
            districts = json.loads(Path(a.districts).read_text(encoding="utf-8"))
        except (OSError, ValueError) as e:
            raise Refused(f"the district table ({a.districts}) will not read, and every race is "
                          f"held to it: {e}")
        doc = read_all(".", a.year, districts, a.saved)
    except Refused as e:
        print(f"REFUSED: {e}", file=sys.stderr)
        return 1
    if a.check:
        try:
            committed = json.loads(out.read_text(encoding="utf-8"))
        except (OSError, ValueError) as e:
            print(f"REFUSED: {out.as_posix()} will not read: {e}", file=sys.stderr)
            return 1
        diff = differences(doc, committed)
        if diff:
            print(f"{out.as_posix()} is not what the saved files say ({len(diff)}):",
                  file=sys.stderr)
            for d in diff[:20]:
                print(f"  {d}", file=sys.stderr)
            return 1
        print(f"{out.as_posix()} is what the saved files say. " + summary(doc))
        return 0
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(dumps(doc), encoding="utf-8")
    for n in doc.get("_notes") or []:
        print(f"FLAGGED: {n}")
    print(f"wrote {out.as_posix()}: " + summary(doc))
    print("  The files mark no winner: each race's winners are its top counts, as many as its "
          "seats.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
